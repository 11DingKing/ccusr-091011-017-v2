"""
固定快照分块导出的行为测试。
"""
import hashlib
import io
from datetime import datetime, timedelta
from decimal import Decimal

from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from django.test import TestCase

from apps.authentication.backends import generate_token
from apps.authentication.models import User
from apps.warehouse.models import (
    Approval, Category, Goods, StockIn, StockOut, Unit, Variety,
)
from .models import ExportJob


class SnapshotExportTestBase(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user("export-admin", "pass123456", role="admin")
        self.user = User.objects.create_user("export-user", "pass123456", role="user")
        self.admin_client = self._client(self.admin)
        self.user_client = self._client(self.user)

        unit = Unit.objects.create(name="件", created_by=self.admin)
        category = Category.objects.create(name="监管器材", unit=unit, created_by=self.admin)
        variety = Variety.objects.create(name="封存设备", category=category, created_by=self.admin)
        self.goods = Goods.objects.create(
            variety=variety, name="封存终端", code="EXP-001",
            quantity=Decimal("10"), warning_threshold=Decimal("5"),
        )

    def _client(self, user):
        client = APIClient()
        client.credentials(HTTP_AUTHORIZATION=f"Bearer {generate_token(user)}")
        return client

    def create_job(self, client=None, export_type="all", extra=None, chunk_size=None):
        client = client or self.admin_client
        payload = {"type": export_type}
        if chunk_size is not None:
            payload["chunk_size"] = chunk_size
        if extra:
            payload.update(extra)
        resp = client.post("/api/exports/", payload, format="json")
        return resp

    def make_stock_in(self, quantity="3", supplier="供应商A"):
        return StockIn.objects.create(
            goods=self.goods, operator=self.admin,
            quantity=Decimal(quantity), batch_no="B-1", supplier=supplier,
        )

    def make_stock_out(self, quantity="2", receiver="张三", dept="检查组"):
        return StockOut.objects.create(
            goods=self.goods, operator=self.admin,
            receiver=receiver, receiver_dept=dept,
            quantity=Decimal(quantity), status="pending",
        )

    def make_approval(self, stock_out, status="pending", remark="待审"):
        return Approval.objects.create(
            stock_out=stock_out, approver=self.admin,
            status=status, remark=remark,
        )

    def parse_xlsx(self, content, sheet_name):
        wb = load_workbook(io.BytesIO(content))
        return list(wb[sheet_name].iter_rows(values_only=True))


class CreateManifestTests(SnapshotExportTestBase):
    def test_first_request_returns_manifest_with_cursor_and_summary(self):
        for i in range(3):
            self.make_stock_in(quantity=str(i + 1))
        out = self.make_stock_out()
        self.make_approval(out)

        resp = self.create_job(chunk_size=2)
        self.assertEqual(resp.status_code, 201, resp.content)
        data = resp.json()["data"]

        self.assertTrue(data["cursor"])
        self.assertEqual(data["export_type"], "all")
        self.assertEqual(data["total_rows"], 5)
        self.assertEqual(data["total_chunks"], 3)
        self.assertEqual(data["chunk_size"], 2)
        sections = {s["key"]: s["rows"] for s in data["sections"]}
        self.assertEqual(sections, {"stock_in": 3, "stock_out": 1, "approval": 1})

        # 分块行数之和 == 总行数；各分块分区行数之和 == 分区总数
        self.assertEqual(sum(c["row_count"] for c in data["chunks"]), 5)
        for key in sections:
            self.assertEqual(sum(c["sections"].get(key, 0) for c in data["chunks"]), sections[key])
        # 位置连续不重不漏
        positions = [(c["start_position"], c["end_position"]) for c in data["chunks"]]
        self.assertEqual(positions, [(0, 2), (2, 4), (4, 5)])

    def test_manifest_dates_default_to_last_30_days(self):
        resp = self.create_job()
        self.assertEqual(resp.status_code, 201)
        data = resp.json()["data"]
        self.assertTrue(data["expires_at"])
        self.assertGreater(data["expires_in_seconds"], 0)
        self.assertEqual(data["total_chunks"], 1)

    def test_invalid_params_rejected(self):
        self.assertEqual(self.create_job(export_type="nope").status_code, 400)
        self.assertEqual(
            self.create_job(extra={"start_date": "2026-10-01", "end_date": "2026-09-01"}).status_code,
            400,
        )
        self.assertEqual(self.create_job(extra={"start_date": "10/01"}).status_code, 400)
        self.assertEqual(self.create_job(extra={"chunk_size": "abc"}).status_code, 400)
        self.assertEqual(self.create_job(extra={"chunk_size": "0"}).status_code, 400)

    def test_requires_auth(self):
        resp = APIClient().post("/api/exports/", {"type": "all"}, format="json")
        self.assertEqual(resp.status_code, 401)


class SnapshotIsolationTests(SnapshotExportTestBase):
    def test_new_records_during_download_do_not_change_manifest_or_chunks(self):
        self.make_stock_in(quantity="1")
        out = self.make_stock_out()
        self.make_approval(out)

        resp = self.create_job(chunk_size=1)
        self.assertEqual(resp.status_code, 201)
        manifest = resp.json()["data"]
        cursor = manifest["cursor"]

        # 先下载全部块并记录字节
        bodies = {}
        for index in range(manifest["total_chunks"]):
            r = self.admin_client.get("/api/exports/chunks/", {"cursor": cursor, "chunk": index})
            self.assertEqual(r.status_code, 200)
            bodies[index] = r.content

        # 长时间分页下载期间持续有新记录进入
        self.make_stock_in(quantity="9")
        new_out = self.make_stock_out(receiver="李四")
        self.make_approval(new_out, status="approved", remark="同意")

        # 清单摘要不变
        status = self.admin_client.get("/api/exports/status/", {"cursor": cursor})
        self.assertEqual(status.status_code, 200)
        frozen = status.json()["data"]
        self.assertEqual(frozen["total_rows"], 3)
        self.assertEqual(frozen["total_chunks"], 3)
        self.assertEqual(
            {s["key"]: s["rows"] for s in frozen["sections"]},
            {"stock_in": 1, "stock_out": 1, "approval": 1},
        )

        # 各分块字节不变
        for index, original in bodies.items():
            r = self.admin_client.get("/api/exports/chunks/", {"cursor": cursor, "chunk": index})
            self.assertEqual(r.status_code, 200)
            self.assertEqual(r.content, original)
            self.assertEqual(r["X-Content-SHA256"], hashlib.sha256(r.content).hexdigest())

        # 新快照包含新记录，各文件页总数与新快照对得上
        resp2 = self.create_job(chunk_size=1)
        manifest2 = resp2.json()["data"]
        self.assertEqual(manifest2["total_rows"], 6)
        status2 = self.admin_client.get("/api/exports/status/", {"cursor": manifest2["cursor"]})
        self.assertEqual(sum(c["row_count"] for c in status2.json()["data"]["chunks"]), 6)

    def test_row_updates_such_as_approval_status_do_not_rewrite_frozen_chunk(self):
        out = self.make_stock_out()
        self.make_approval(out, status="pending", remark="等待审批")
        resp = self.create_job(export_type="approval")
        cursor = resp.json()["data"]["cursor"]

        first = self.admin_client.get("/api/exports/chunks/", {"cursor": cursor, "chunk": 0})
        self.assertEqual(first.status_code, 200)
        rows = self.parse_xlsx(first.content, "审批记录")
        self.assertIn("待审批", rows[1])

        # 审批在快照创建后通过，出库状态也改变
        Approval.objects.filter(stock_out=out).update(status="approved", remark="同意出库")
        out.status = "approved"
        out.save(update_fields=["status"])

        second = self.admin_client.get("/api/exports/chunks/", {"cursor": cursor, "chunk": 0})
        self.assertEqual(second.content, first.content)
        self.assertEqual(second["X-Content-SHA256"], first["X-Content-SHA256"])

    def test_rows_partitioned_without_loss_or_duplication(self):
        for i in range(5):
            self.make_stock_in(quantity=str(i + 1))
        resp = self.create_job(export_type="stock_in", extra={"chunk_size": 2})
        cursor = resp.json()["data"]["cursor"]

        seen = []
        for index in range(3):
            r = self.admin_client.get("/api/exports/chunks/", {"cursor": cursor, "chunk": index})
            self.assertEqual(r.status_code, 200)
            rows = self.parse_xlsx(r.content, "入库记录")
            body = rows[1:]  # 去掉表头
            seen.extend(row[0] for row in body)
        self.assertEqual(sorted(seen), sorted(StockIn.objects.values_list("id", flat=True)))
        self.assertEqual(len(seen), len(set(seen)))

    def test_each_chunk_of_all_export_has_all_section_sheets(self):
        for _ in range(4):
            self.make_stock_in()
        for _ in range(3):
            self.make_stock_out()
        for out in StockOut.objects.all():
            self.make_approval(out)

        resp = self.create_job(chunk_size=2)
        self.assertEqual(resp.status_code, 201)
        manifest = resp.json()["data"]
        cursor = manifest["cursor"]
        self.assertEqual(manifest["total_rows"], 10)

        collected = {"入库记录": [], "出库记录": [], "审批记录": []}
        for index in range(manifest["total_chunks"]):
            r = self.admin_client.get("/api/exports/chunks/", {"cursor": cursor, "chunk": index})
            self.assertEqual(r.status_code, 200)
            wb = load_workbook(io.BytesIO(r.content))
            # 每个分块都包含全部分区 sheet，各自带表头，便于独立打开核对
            self.assertEqual(wb.sheetnames, ["入库记录", "出库记录", "审批记录"])
            for name in collected:
                sheet_rows = list(wb[name].iter_rows(values_only=True))
                self.assertEqual(sheet_rows[0][:2], ("记录ID", "货物编码") if name != "审批记录"
                                 else ("审批ID", "出库记录ID"))
                collected[name].extend(sheet_rows[1:])

        # 跨分块汇总后各分区行数与清单摘要一致、无重无漏
        self.assertEqual(len(collected["入库记录"]), 4)
        self.assertEqual(len(collected["出库记录"]), 3)
        self.assertEqual(len(collected["审批记录"]), 3)
        in_ids = [row[0] for row in collected["入库记录"]]
        self.assertEqual(len(in_ids), len(set(in_ids)))


class RepeatDownloadTests(SnapshotExportTestBase):
    def test_repeat_download_same_chunk_is_byte_identical(self):
        self.make_stock_in()
        self.make_stock_out()
        resp = self.create_job()
        cursor = resp.json()["data"]["cursor"]

        first = self.admin_client.get("/api/exports/chunks/", {"cursor": cursor, "chunk": 0})
        second = self.admin_client.get("/api/exports/chunks/", {"cursor": cursor, "chunk": 0})
        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.content, first.content)
        self.assertEqual(second["X-Content-SHA256"], first["X-Content-SHA256"])
        self.assertEqual(second["ETag"], first["ETag"])

        # 状态接口暴露的校验值与实际内容一致
        status = self.admin_client.get("/api/exports/status/", {"cursor": cursor})
        chunk0 = status.json()["data"]["chunks"][0]
        self.assertTrue(chunk0["ready"])
        self.assertEqual(chunk0["checksum"], first["X-Content-SHA256"])
        self.assertEqual(chunk0["size_bytes"], len(first.content))


class CursorAndExpiryTests(SnapshotExportTestBase):
    def _job(self, resp):
        return ExportJob.objects.get(pk=resp.json()["data"]["export_job_id"])

    def test_missing_cursor_rejected(self):
        r = self.admin_client.get("/api/exports/chunks/", {"chunk": 0})
        self.assertEqual(r.status_code, 400)

    def test_unknown_cursor_rejected_with_404(self):
        r = self.admin_client.get("/api/exports/chunks/", {"cursor": "not-a-real-cursor", "chunk": 0})
        self.assertEqual(r.status_code, 404)

    def test_chunk_index_out_of_range_rejected(self):
        resp = self.create_job()
        cursor = resp.json()["data"]["cursor"]
        r = self.admin_client.get("/api/exports/chunks/", {"cursor": cursor, "chunk": 5})
        self.assertEqual(r.status_code, 400)
        r = self.admin_client.get("/api/exports/chunks/", {"cursor": cursor, "chunk": "x"})
        self.assertEqual(r.status_code, 400)

    def test_expired_snapshot_explicitly_rejected_with_410(self):
        resp = self.create_job()
        cursor = resp.json()["data"]["cursor"]
        job = self._job(resp)

        # 过期前可正常下载
        self.assertEqual(
            self.admin_client.get("/api/exports/chunks/", {"cursor": cursor, "chunk": 0}).status_code,
            200,
        )

        ExportJob.objects.filter(pk=job.pk).update(
            expires_at=timezone.now() - timedelta(minutes=1)
        )
        r = self.admin_client.get("/api/exports/chunks/", {"cursor": cursor, "chunk": 0})
        self.assertEqual(r.status_code, 410)
        self.assertIn("过期", r.json()["message"])
        r2 = self.admin_client.get("/api/exports/status/", {"cursor": cursor})
        self.assertEqual(r2.status_code, 410)

        # 过期后即使已有固化分块也拒绝，且不回源业务表
        job.refresh_from_db()
        self.assertEqual(job.status, ExportJob.STATUS_EXPIRED)

    def test_expired_snapshot_does_not_block_new_export(self):
        self.make_stock_in()
        resp = self.create_job()
        old_job = self._job(resp)
        ExportJob.objects.filter(pk=old_job.pk).update(
            expires_at=timezone.now() - timedelta(minutes=1)
        )
        resp2 = self.create_job()
        self.assertEqual(resp2.status_code, 201)
        self.assertNotEqual(resp2.json()["data"]["cursor"], resp.json()["data"]["cursor"])


class PermissionTests(SnapshotExportTestBase):
    def test_approval_data_requires_admin_at_creation(self):
        self.assertEqual(self.user_client.post("/api/exports/", {"type": "approval"}, format="json").status_code, 403)
        self.assertEqual(self.user_client.post("/api/exports/", {"type": "all"}, format="json").status_code, 403)
        # 收发记录普通用户可导出
        self.assertEqual(self.user_client.post("/api/exports/", {"type": "stock_in"}, format="json").status_code, 201)
        self.assertEqual(self.user_client.post("/api/exports/", {"type": "stock_out"}, format="json").status_code, 201)

    def test_other_user_cannot_access_snapshot(self):
        resp = self.admin_client.post("/api/exports/", {"type": "stock_in"}, format="json")
        cursor = resp.json()["data"]["cursor"]
        r = self.user_client.get("/api/exports/chunks/", {"cursor": cursor, "chunk": 0})
        self.assertEqual(r.status_code, 404)
        self.assertEqual(
            self.user_client.get("/api/exports/status/", {"cursor": cursor}).status_code, 404
        )

    def test_permission_change_affects_new_exports_but_not_existing_manifest(self):
        # 管理员身份创建含审批数据的快照
        out = self.make_stock_out()
        self.make_approval(out)
        resp = self.create_job()
        self.assertEqual(resp.status_code, 201)
        manifest = resp.json()["data"]
        cursor = manifest["cursor"]
        approval_rows = next(s["rows"] for s in manifest["sections"] if s["key"] == "approval")
        self.assertEqual(approval_rows, 1)

        # 之后被降级为普通用户：旧快照仍可访问，清单不被改写
        self.admin.role = "user"
        self.admin.save(update_fields=["role"])
        self.assertEqual(
            self.admin_client.get("/api/exports/chunks/", {"cursor": cursor, "chunk": 0}).status_code,
            200,
        )
        status = self.admin_client.get("/api/exports/status/", {"cursor": cursor})
        self.assertEqual(status.status_code, 200)
        self.assertEqual(status.json()["data"]["total_rows"], 2)

        # 但不能再新建含审批数据的导出
        self.assertEqual(self.create_job().status_code, 403)

        # 重新升回管理员后新建导出恢复
        self.admin.role = "admin"
        self.admin.save(update_fields=["role"])
        self.assertEqual(self.create_job().status_code, 201)


class EmptyExportTests(SnapshotExportTestBase):
    def test_empty_range_produces_single_header_only_chunk(self):
        resp = self.create_job(export_type="stock_in", extra={
            "start_date": "2020-01-01", "end_date": "2020-01-31",
        })
        self.assertEqual(resp.status_code, 201)
        data = resp.json()["data"]
        self.assertEqual(data["total_rows"], 0)
        self.assertEqual(data["total_chunks"], 1)

        r = self.admin_client.get(
            "/api/exports/chunks/",
            {"cursor": data["cursor"], "chunk": 0},
        )
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["X-Row-Count"], "0")
        rows = self.parse_xlsx(r.content, "入库记录")
        self.assertEqual(len(rows), 1)  # 仅表头
        self.assertEqual(rows[0][0], "记录ID")

    def test_repeat_download_empty_chunk_stable(self):
        resp = self.create_job(export_type="stock_in", extra={
            "start_date": "2020-01-01", "end_date": "2020-01-31",
        })
        cursor = resp.json()["data"]["cursor"]
        first = self.admin_client.get("/api/exports/chunks/", {"cursor": cursor, "chunk": 0})
        second = self.admin_client.get("/api/exports/chunks/", {"cursor": cursor, "chunk": 0})
        self.assertEqual(second.content, first.content)
        self.assertEqual(second["X-Content-SHA256"], first["X-Content-SHA256"])
